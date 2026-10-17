"""Shared-folder artifacts: append CSV/XLSX and upload via presigned URL."""
from __future__ import annotations

import csv
import os
from pathlib import Path

import httpx

from .policy import in_roots


def append_rows(path: str, rows: list[dict], roots: list[str], fmt: str = "csv") -> dict:
    real = os.path.realpath(os.path.expanduser(path))
    if not in_roots(real, roots):
        return {"status": "error", "reason": "path_outside_shared_folders"}
    Path(real).parent.mkdir(parents=True, exist_ok=True)
    if fmt == "csv" or real.lower().endswith(".csv"):
        write_header = not os.path.isfile(real) or os.path.getsize(real) == 0
        with open(real, "a", newline="", encoding="utf-8") as fh:
            if not rows:
                return {"status": "ok", "path": real, "appended": 0}
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            if write_header:
                w.writeheader()
            w.writerows(rows)
        return {"status": "ok", "path": real, "appended": len(rows)}
    if fmt == "xlsx" or real.lower().endswith(".xlsx"):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            return {"status": "error", "reason": "openpyxl_not_installed"}
        if os.path.isfile(real):
            wb = load_workbook(real)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            if rows:
                ws.append(list(rows[0].keys()))
        for row in rows:
            ws.append([row.get(k) for k in row.keys()])
        wb.save(real)
        return {"status": "ok", "path": real, "appended": len(rows)}
    return {"status": "error", "reason": "unsupported_format"}


def upload_presigned(url: str, path: str, roots: list[str], content_type: str | None = None) -> dict:
    real = os.path.realpath(os.path.expanduser(path))
    if not in_roots(real, roots):
        return {"status": "error", "reason": "path_outside_shared_folders"}
    with open(real, "rb") as fh:
        data = fh.read()
    headers = {"Content-Type": content_type or "application/octet-stream"}
    r = httpx.put(url, content=data, headers=headers, timeout=120.0)
    if r.status_code >= 300:
        return {"status": "error", "reason": f"upload_http_{r.status_code}"}
    return {"status": "ok", "bytes": len(data), "upload_url": url.split("?")[0]}
